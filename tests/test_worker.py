import asyncio
from datetime import datetime

import pytest
from openpyxl import load_workbook

from foodbot import claude_runner, messages
from foodbot.queue_worker import Job, Worker
from foodbot.records import Records
from foodbot.schema import AnalysisError
from foodbot.usage import Usage

JPG = b"\xff\xd8\xff\xe0" + b"0" * 20
FOOD = {"is_food": True, "name": "雞腿便當", "unclear": False,
        "ingredients": [{"name": "白飯", "weight_g": 150, "kcal": 200},
                        {"name": "雞腿", "weight_g": 120, "kcal": 250}]}
MUG = {"is_food": False, "name": "馬克杯", "unclear": False, "ingredients": []}


class Harness:
    def __init__(self, cfg):
        self.usage = Usage(cfg)
        self.records = Records(cfg)
        self.worker = Worker(cfg, self.usage, self.records)
        self.texts: dict[int, list[str]] = {}
        self.finished: list[int] = []

    def job(self, n, uid=1, data=JPG, size=None, when=None):
        self.texts[n] = []

        async def set_text(t):
            self.texts[n].append(t)

        async def finished():
            self.finished.append(n)

        async def download():
            return data

        j = Job(uid, "小明", when or datetime(2026, 10, 20, 12, 0, n), "附帶文字",
                len(data) if size is None else size, download, set_text, finished)
        self.worker.submit(j)
        return j

    async def drain(self, count):
        task = asyncio.create_task(self.worker.run())
        for _ in range(500):
            if len(self.finished) >= count:
                break
            await asyncio.sleep(0.01)
        task.cancel()


def fake_analyze(answers):
    answers = list(answers)

    async def analyze(cfg, photo, ext, hint):
        a = answers.pop(0)
        if isinstance(a, Exception):
            raise a
        return a
    return analyze


async def test_food_then_not_food_flow(cfg, monkeypatch):
    monkeypatch.setattr(claude_runner, "analyze", fake_analyze([FOOD, MUG]))
    h = Harness(cfg)
    h.job(1)
    h.job(2)
    await h.drain(2)
    assert "你今天已攝取：450 kcal" in h.texts[1][-1] and "<@1>" in h.texts[1][-1]
    assert "是馬克杯" in h.texts[2][-1] and "450 kcal（沒有變動）" in h.texts[2][-1]
    ws = load_workbook(cfg.records_path).active
    assert ws.max_row == 3 and ws.cell(3, 4).value == "不是食物" and ws.cell(3, 5).value == "馬克杯"
    assert len(list(cfg.photos_dir.iterdir())) == 2
    assert h.usage.used_total() == 2


async def test_failures_do_not_record_but_count_and_do_not_stop_worker(cfg, monkeypatch):
    bad = [AnalysisError("逾時"), {"is_food": True, "name": "x", "unclear": False,
                                   "ingredients": [{"name": "a", "weight_g": 1, "kcal": "很多"}]}, FOOD]
    monkeypatch.setattr(claude_runner, "analyze", fake_analyze(bad))
    h = Harness(cfg)
    for n in (1, 2, 3):
        h.job(n)
    await h.drain(3)
    assert h.texts[1][-1] == messages.FAILED and h.texts[2][-1] == messages.FAILED
    assert "450" in h.texts[3][-1]  # 失敗的沒累加
    assert load_workbook(cfg.records_path).active.max_row == 2
    assert h.usage.used_total() == 3  # 失敗也算額度


async def test_bad_image_fails_before_claude_and_is_not_counted(cfg, monkeypatch):
    async def boom(*a):
        raise AssertionError("不該呼叫 claude")
    monkeypatch.setattr(claude_runner, "analyze", boom)
    h = Harness(cfg)
    h.job(1, data=b"%PDF-1.4 not an image")
    h.job(2, size=11 * 1024 * 1024)
    await h.drain(2)
    assert h.texts[1][-1] == messages.FAILED and h.texts[2][-1] == messages.FAILED
    assert h.usage.used_total() == 0 and not cfg.records_path.exists()


async def test_queue_positions_and_bookkeeping(cfg, monkeypatch):
    monkeypatch.setattr(claude_runner, "analyze", fake_analyze([FOOD, FOOD, FOOD]))
    h = Harness(cfg)
    for n in (1, 2, 3):
        h.job(n, uid=n)
    assert h.worker.queued_total() == 3 and h.worker.queued_user(2) == 1
    assert h.worker.free_slots() == 7
    await h.drain(3)
    assert h.texts[3][0] == messages.processing(2)  # 第 1 張開始時，3 號前面還有 2 張
    assert h.texts[3][1] == messages.processing(1)
    assert h.texts[3][2] == messages.processing(0)
    assert h.finished == [1, 2, 3]
    assert h.worker.queued_total() == 0


async def test_per_user_totals_are_separate(cfg, monkeypatch):
    monkeypatch.setattr(claude_runner, "analyze", fake_analyze([FOOD, FOOD, FOOD]))
    h = Harness(cfg)
    h.job(1, uid=1)
    h.job(2, uid=2)
    h.job(3, uid=1)
    await h.drain(3)
    assert "450 kcal" in h.texts[1][-1] and "450 kcal" in h.texts[2][-1]
    assert "900 kcal" in h.texts[3][-1]


def test_no_everyone_ping_in_reply():
    from foodbot.schema import validate
    a = validate({**FOOD, "name": "@everyone 便當"})
    assert "@everyone" not in messages.result(1, a, 100)

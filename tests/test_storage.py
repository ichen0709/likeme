import json
import os
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from foodbot import photos, records as records_mod
from foodbot.records import Record, Records
from foodbot.usage import Usage

JPG = b"\xff\xd8\xff\xe0" + b"0" * 20


def rec(when, uid="1", kcal=100.0, total=100.0, name="雞腿便當", uploader="小明"):
    return Record(when, uploader, uid, name, "白飯 150g 200kcal", 150, kcal, total, "C:\\p.jpg")


# ---- photos ----
def test_sniff_image():
    assert photos.sniff_image(JPG) == "jpg"
    assert photos.sniff_image(b"\x89PNG\r\n\x1a\n....") == "png"
    assert photos.sniff_image(b"RIFF1234WEBPxxxx") == "webp"
    assert photos.sniff_image(b"GIF89a....") == "gif"
    assert photos.sniff_image(b"%PDF-1.4") is None
    assert photos.sniff_image(b"") is None


def test_photo_filename_format_and_collisions(tmp_path):
    when = datetime(2026, 10, 20, 12, 34, 56)
    p1 = photos.save_photo(tmp_path, when, "白飯", JPG, "jpg")
    p2 = photos.save_photo(tmp_path, when, "白飯", JPG + b"2", "jpg")
    p3 = photos.save_photo(tmp_path, when, "白飯", JPG + b"3", "jpg")
    assert p1.name == "2026-10-20 12-34-56 白飯.jpg"
    assert p2.name == "2026-10-20 12-34-56 白飯-2.jpg"
    assert p3.name == "2026-10-20 12-34-56 白飯-3.jpg"
    assert p1.read_bytes() == JPG and p2.read_bytes() == JPG + b"2"


def test_photo_evil_name_stays_inside(tmp_path):
    photos_dir = tmp_path / "photos"
    p = photos.save_photo(photos_dir, datetime(2026, 10, 20, 1, 2, 3), "..\\..\\Windows\\x", JPG, "jpg")
    assert p.parent == photos_dir.resolve() or p.parent == photos_dir
    assert list(tmp_path.iterdir()) == [photos_dir]


# ---- usage ----
def test_usage_limits_and_persistence(cfg):
    u = Usage(cfg)
    assert u.remaining(1, 0, 0) == (15, False)
    for _ in range(13):
        u.increment(1)
    assert u.remaining(1, 0, 0) == (2, False)
    assert u.remaining(1, 1, 1) == (1, False)  # 排隊中的也算
    assert u.remaining(2, 0, 0) == (15, False)  # 別人不受影響
    assert Usage(cfg).used_user(1) == 13  # 重開還在


def test_usage_total_limit(cfg, monkeypatch):
    monkeypatch.setattr("foodbot.usage.DAILY_LIMIT_TOTAL", 3)
    u = Usage(cfg)
    for uid in (1, 2, 3):
        u.increment(uid)
    assert u.remaining(4, 0, 0) == (0, True)


def test_usage_corrupt_file_resets(cfg):
    cfg.usage_path.write_text("{not json", encoding="utf-8")
    assert Usage(cfg).used_total() == 0


def test_usage_new_day_resets(cfg):
    cfg.usage_path.write_text(json.dumps({"date": "2000-01-01", "total": 99, "users": {"1": 15}}))
    assert Usage(cfg).used_user(1) == 0


# ---- records ----
def test_records_append_never_overwrites(cfg):
    r = Records(cfg)
    when = datetime(2026, 10, 20, 12, 0, 0)
    assert r.append(rec(when))
    assert Records(cfg).append(rec(when, uid="2", total=50, kcal=50))  # 重開後再加
    ws = load_workbook(cfg.records_path).active
    assert ws.max_row == 3
    assert [c.value for c in ws[1]][:3] == ["日期時間", "上傳者", "使用者ID"]
    assert ws.cell(2, 3).value == "1" and ws.cell(3, 3).value == "2"
    assert isinstance(ws.cell(2, 7).value, (int, float)) and ws.cell(2, 1).value == when


def test_daily_total_per_user_per_day_and_manual_edit(cfg):
    r = Records(cfg)
    d1, d2 = datetime(2026, 10, 20, 23, 59), datetime(2026, 10, 21, 0, 1)
    r.append(rec(d1, "1", 100, 100))
    r.append(rec(d1, "1", 200, 300))
    r.append(rec(d1, "2", 999, 999))
    r.append(rec(d2, "1", 50, 50))
    assert r.daily_total("1", d1.date()) == 300
    assert r.daily_total("2", d1.date()) == 999
    assert r.daily_total("1", d2.date()) == 50
    wb = load_workbook(cfg.records_path)
    wb.active.cell(2, 7).value = 500  # 手動改熱量
    wb.save(cfg.records_path)
    assert r.daily_total("1", d1.date()) == 700


def test_not_food_row_has_blank_numbers_and_zero_effect(cfg):
    r = Records(cfg)
    when = datetime(2026, 10, 20, 12)
    r.append(rec(when, kcal=100, total=100))
    r.append(Record(when, "小明", "1", "不是食物", "馬克杯", None, None, 100, "p"))
    ws = load_workbook(cfg.records_path).active
    assert ws.cell(3, 6).value is None and ws.cell(3, 7).value is None
    assert ws.cell(3, 4).value == "不是食物" and ws.cell(3, 5).value == "馬克杯"
    assert r.daily_total("1", when.date()) == 100


def test_formula_text_in_excel_is_text(cfg):
    r = Records(cfg)
    r.append(rec(datetime(2026, 10, 20, 12), name="=1+1", uploader="@everyone"))
    ws = load_workbook(cfg.records_path).active
    assert ws.cell(2, 4).value == "'=1+1" and ws.cell(2, 4).data_type == "s"
    assert ws.cell(2, 2).value == "'@everyone"


def test_excel_locked_goes_to_pending_then_backfills_in_order(cfg, monkeypatch):
    monkeypatch.setattr(records_mod, "RETRY_DELAY_SECONDS", 0)
    r = Records(cfg)
    when = datetime(2026, 10, 20, 12)
    assert r.append(rec(when, kcal=100, total=100, name="第一"))
    real_replace = os.replace

    def locked(src, dst):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(records_mod.os, "replace", locked)
    assert not r.append(rec(when, kcal=200, total=300, name="第二"))
    assert cfg.pending_path.exists()
    assert r.daily_total("1", when.date()) == 300  # Excel + 暫存都算
    monkeypatch.setattr(records_mod.os, "replace", real_replace)
    assert r.append(rec(when, kcal=50, total=350, name="第三"))
    assert not cfg.pending_path.exists()
    ws = load_workbook(cfg.records_path).active
    assert [ws.cell(i, 4).value for i in (2, 3, 4)] == ["第一", "第二", "第三"]


def test_daily_total_falls_back_to_cache_when_excel_unreadable(cfg, monkeypatch):
    r = Records(cfg)
    when = datetime(2026, 10, 20, 12)
    r.append(rec(when, kcal=100, total=100))
    monkeypatch.setattr(r, "_xlsx_total", lambda *a: (_ for _ in ()).throw(PermissionError()))
    assert r.daily_total("1", date(2026, 10, 20)) == 100

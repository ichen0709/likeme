"""Excel 紀錄：往下加列、Excel 被占用時暫存、當日總數（架構文件第 8 節）。"""

import json
import os
import time
from dataclasses import asdict, dataclass
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from . import guards
from .config import Config

HEADERS = ["日期時間", "上傳者", "使用者ID", "食物敘述", "成份", "食物重量", "食物熱量",
           "當日已攝取熱量", "照片存放位置"]
WRITE_ATTEMPTS = 3
RETRY_DELAY_SECONDS = 1.0


@dataclass
class Record:
    when: datetime
    uploader: str
    user_id: str
    description: str
    ingredients: str
    weight_g: float | None
    kcal: float | None
    day_total: float
    photo_path: str


class Records:
    def __init__(self, cfg: Config):
        self._path = cfg.records_path
        self._pending = cfg.pending_path
        self._last_xlsx_total: dict[tuple[str, date], float] = {}

    # ---- 讀 ----
    def _pending_rows(self) -> list[dict]:
        try:
            lines = self._pending.read_text(encoding="utf-8").splitlines()
        except FileNotFoundError:
            return []
        rows = []
        for line in lines:
            try:
                rows.append(json.loads(line))
            except ValueError:
                print("警告：pending.jsonl 有一行讀不懂，已略過。")
        return rows

    def _pending_total(self, user_id: str, day: date) -> float:
        total = 0.0
        for row in self._pending_rows():
            if row.get("user_id") == user_id and row.get("when", "")[:10] == day.isoformat():
                total += row.get("kcal") or 0
        return total

    def _xlsx_total(self, user_id: str, day: date) -> float:
        if not self._path.exists():
            return 0.0
        wb = load_workbook(self._path, read_only=True, data_only=True)
        try:
            total = 0.0
            for row in wb.active.iter_rows(min_row=2, max_col=7, values_only=True):
                when, uid, kcal = row[0], row[2], row[6]
                if uid is None or str(uid).lstrip("'") != user_id:
                    continue
                if isinstance(when, str):
                    try:
                        when = datetime.fromisoformat(when)
                    except ValueError:
                        continue
                if (isinstance(when, datetime) and when.date() == day
                        and isinstance(kcal, (int, float)) and not isinstance(kcal, bool)):
                    total += kcal
            return total
        finally:
            wb.close()

    def daily_total(self, user_id: str, day: date) -> float:
        """該使用者該日已記錄的熱量（不含這一餐）：Excel 加暫存。"""
        pending = self._pending_total(user_id, day)
        try:
            xlsx = self._xlsx_total(user_id, day)
        except Exception as e:
            cached = self._last_xlsx_total.get((user_id, day), 0.0)
            print(f"警告：讀不到 records.xlsx（{type(e).__name__}），"
                  f"改用上次讀到的總數 {cached:g} 估算。")
            return cached + pending
        self._last_xlsx_total[(user_id, day)] = xlsx
        return xlsx + pending

    # ---- 寫 ----
    def _write_rows(self, rows: list[Record]) -> None:
        if self._path.exists():
            wb = load_workbook(self._path)
        else:
            wb = Workbook()
            wb.active.append(HEADERS)
        ws = wb.active
        for r in rows:
            ws.append([r.when, None, None, None, None, r.weight_g, r.kcal, r.day_total, None])
            row = ws.max_row
            ws.cell(row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
            guards.write_text_cell(ws.cell(row, 2), r.uploader)
            guards.write_text_cell(ws.cell(row, 3), r.user_id)
            guards.write_text_cell(ws.cell(row, 4), r.description)
            guards.write_text_cell(ws.cell(row, 5), r.ingredients)
            guards.write_text_cell(ws.cell(row, 9), r.photo_path)
        tmp = self._path.with_name(self._path.name + ".tmp")
        wb.save(tmp)
        os.replace(tmp, self._path)

    def append(self, record: Record) -> bool:
        """加一列到 Excel（會先補寫暫存的）。被占用就重試，仍失敗就存進 pending.jsonl。
        回傳是否已寫進 Excel。"""
        backlog = [Record(**{**row, "when": datetime.fromisoformat(row["when"])})
                   for row in self._pending_rows()]
        rows = backlog + [record]
        for attempt in range(1, WRITE_ATTEMPTS + 1):
            try:
                self._write_rows(rows)
            except OSError as e:
                print(f"警告：records.xlsx 寫入失敗（{e.strerror}），可能正被 Excel 開著"
                      f"（第 {attempt}/{WRITE_ATTEMPTS} 次）。")
                if attempt < WRITE_ATTEMPTS:
                    time.sleep(RETRY_DELAY_SECONDS)
            else:
                self._pending.unlink(missing_ok=True)
                self._last_xlsx_total[(record.user_id, record.when.date())] = record.day_total
                return True
        self._append_pending(record)
        print("警告：這一筆已暫存到 pending.jsonl，下次寫入成功時會依序補進 Excel。")
        return False

    def _append_pending(self, record: Record) -> None:
        row = asdict(record)
        row["when"] = record.when.isoformat()
        with open(self._pending, "a", encoding="utf-8") as f:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
